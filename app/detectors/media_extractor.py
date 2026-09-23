"""
Media Content Extractor for PromptGuard AI Firewall Gateway.

Extracts readable text from images (via OCR) and document files (PDF, DOCX, XLSX, CSV, TXT, etc.)
so the existing 4-stage detection pipeline can scan media for sensitive data.

Works with ANY LLM — extraction is done server-side before forwarding.
"""

import base64
import io
import re
import logging
from typing import List, Optional, Tuple
from dataclasses import dataclass, field

logger = logging.getLogger("promptguard.media")

# ── Optional imports (graceful degradation) ──────────────────────────────────

_TESSERACT_AVAILABLE = False
_PIL_AVAILABLE = False
_PYPDF2_AVAILABLE = False
_DOCX_AVAILABLE = False
_OPENPYXL_AVAILABLE = False

try:
    from PIL import Image
    _PIL_AVAILABLE = True
except ImportError:
    logger.warning("Pillow not installed — image processing disabled")

try:
    import pytesseract
    _TESSERACT_AVAILABLE = True
except ImportError:
    logger.warning("pytesseract not installed — image OCR disabled")

try:
    from PyPDF2 import PdfReader
    _PYPDF2_AVAILABLE = True
except ImportError:
    logger.warning("PyPDF2 not installed — PDF extraction disabled")

try:
    import docx
    _DOCX_AVAILABLE = True
except ImportError:
    logger.warning("python-docx not installed — DOCX extraction disabled")

try:
    import openpyxl
    _OPENPYXL_AVAILABLE = True
except ImportError:
    logger.warning("openpyxl not installed — XLSX extraction disabled")


# ── Supported MIME types ─────────────────────────────────────────────────────

IMAGE_MIME_TYPES = {"image/png", "image/jpeg", "image/jpg", "image/gif", "image/bmp", "image/tiff", "image/webp"}

DOCUMENT_MIME_TYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",  # .docx
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",         # .xlsx
    "text/csv",
    "text/plain",
    "application/json",
    "application/xml",
    "text/xml",
    "text/yaml",
    "text/x-yaml",
    "application/x-yaml",
}

# Extension → MIME fallback map
EXTENSION_MIME_MAP = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".bmp": "image/bmp",
    ".tiff": "image/tiff",
    ".tif": "image/tiff",
    ".webp": "image/webp",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".csv": "text/csv",
    ".txt": "text/plain",
    ".log": "text/plain",
    ".env": "text/plain",
    ".ini": "text/plain",
    ".conf": "text/plain",
    ".cfg": "text/plain",
    ".json": "application/json",
    ".xml": "application/xml",
    ".yaml": "text/yaml",
    ".yml": "text/yaml",
    ".md": "text/plain",
    ".py": "text/plain",
    ".js": "text/plain",
    ".ts": "text/plain",
    ".java": "text/plain",
    ".go": "text/plain",
    ".rs": "text/plain",
    ".rb": "text/plain",
    ".sh": "text/plain",
    ".sql": "text/plain",
    ".html": "text/plain",
    ".css": "text/plain",
}


@dataclass
class MediaExtractionResult:
    """Result of extracting text from a single media item."""
    source_type: str          # "image" or "file"
    filename: str             # Original filename or "inline_image"
    mime_type: str             # Detected MIME type
    extracted_text: str        # All text extracted from this media
    char_count: int = 0        # Number of characters extracted
    success: bool = True       # Whether extraction succeeded
    error: str = ""            # Error message if failed
    method: str = ""           # Extraction method used (e.g., "tesseract_ocr", "pypdf2", "direct_read")


@dataclass
class MediaScanSummary:
    """Summary of all media extraction performed on a request."""
    total_media_items: int = 0
    images_scanned: int = 0
    files_scanned: int = 0
    total_text_extracted: int = 0        # total characters
    extractions: List[MediaExtractionResult] = field(default_factory=list)
    combined_text: str = ""              # All extracted text concatenated


def _resolve_mime_type(filename: str, provided_mime: Optional[str] = None) -> str:
    """Resolve MIME type from filename extension if not provided."""
    if provided_mime and provided_mime != "application/octet-stream":
        return provided_mime
    if filename:
        ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        return EXTENSION_MIME_MAP.get(ext, "application/octet-stream")
    return "application/octet-stream"


def _decode_base64_content(content_base64: str) -> bytes:
    """Decode base64 content, handling data URI prefix."""
    # Strip data URI prefix if present (e.g., "data:image/png;base64,...")
    if "," in content_base64 and content_base64.startswith("data:"):
        content_base64 = content_base64.split(",", 1)[1]
    return base64.b64decode(content_base64)


# ── Image OCR ────────────────────────────────────────────────────────────────

def extract_text_from_image(image_bytes: bytes, filename: str = "image") -> MediaExtractionResult:
    """Extract text from an image using Tesseract OCR."""
    if not _PIL_AVAILABLE:
        return MediaExtractionResult(
            source_type="image",
            filename=filename,
            mime_type="image/*",
            extracted_text="",
            success=False,
            error="Pillow not installed — cannot process images",
            method="none"
        )

    if not _TESSERACT_AVAILABLE:
        return MediaExtractionResult(
            source_type="image",
            filename=filename,
            mime_type="image/*",
            extracted_text="",
            success=False,
            error="pytesseract not installed — OCR unavailable",
            method="none"
        )

    try:
        image = Image.open(io.BytesIO(image_bytes))
        # Convert to RGB if needed (handles RGBA, palette mode, etc.)
        if image.mode not in ("L", "RGB"):
            image = image.convert("RGB")

        text = pytesseract.image_to_string(image, lang="eng")
        text = text.strip()

        return MediaExtractionResult(
            source_type="image",
            filename=filename,
            mime_type="image/*",
            extracted_text=text,
            char_count=len(text),
            success=True,
            method="tesseract_ocr"
        )
    except Exception as e:
        logger.error(f"Image OCR failed for {filename}: {e}")
        return MediaExtractionResult(
            source_type="image",
            filename=filename,
            mime_type="image/*",
            extracted_text="",
            success=False,
            error=f"OCR failed: {str(e)}",
            method="tesseract_ocr"
        )


# ── Document Text Extraction ─────────────────────────────────────────────────

def extract_text_from_pdf(file_bytes: bytes, filename: str = "document.pdf") -> MediaExtractionResult:
    """Extract text from a PDF file."""
    if not _PYPDF2_AVAILABLE:
        return MediaExtractionResult(
            source_type="file",
            filename=filename,
            mime_type="application/pdf",
            extracted_text="",
            success=False,
            error="PyPDF2 not installed — PDF extraction disabled",
            method="none"
        )

    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        pages_text = []
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                pages_text.append(page_text)
        text = "\n".join(pages_text).strip()

        return MediaExtractionResult(
            source_type="file",
            filename=filename,
            mime_type="application/pdf",
            extracted_text=text,
            char_count=len(text),
            success=True,
            method="pypdf2"
        )
    except Exception as e:
        logger.error(f"PDF extraction failed for {filename}: {e}")
        return MediaExtractionResult(
            source_type="file",
            filename=filename,
            mime_type="application/pdf",
            extracted_text="",
            success=False,
            error=f"PDF extraction failed: {str(e)}",
            method="pypdf2"
        )


def extract_text_from_docx(file_bytes: bytes, filename: str = "document.docx") -> MediaExtractionResult:
    """Extract text from a DOCX file."""
    if not _DOCX_AVAILABLE:
        return MediaExtractionResult(
            source_type="file",
            filename=filename,
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            extracted_text="",
            success=False,
            error="python-docx not installed — DOCX extraction disabled",
            method="none"
        )

    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        text = "\n".join(paragraphs).strip()

        return MediaExtractionResult(
            source_type="file",
            filename=filename,
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            extracted_text=text,
            char_count=len(text),
            success=True,
            method="python_docx"
        )
    except Exception as e:
        logger.error(f"DOCX extraction failed for {filename}: {e}")
        return MediaExtractionResult(
            source_type="file",
            filename=filename,
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            extracted_text="",
            success=False,
            error=f"DOCX extraction failed: {str(e)}",
            method="python_docx"
        )


def extract_text_from_xlsx(file_bytes: bytes, filename: str = "spreadsheet.xlsx") -> MediaExtractionResult:
    """Extract text from an XLSX file."""
    if not _OPENPYXL_AVAILABLE:
        return MediaExtractionResult(
            source_type="file",
            filename=filename,
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            extracted_text="",
            success=False,
            error="openpyxl not installed — XLSX extraction disabled",
            method="none"
        )

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        all_text = []
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            for row in ws.iter_rows(values_only=True):
                row_values = [str(cell) for cell in row if cell is not None]
                if row_values:
                    all_text.append(" | ".join(row_values))
        wb.close()
        text = "\n".join(all_text).strip()

        return MediaExtractionResult(
            source_type="file",
            filename=filename,
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            extracted_text=text,
            char_count=len(text),
            success=True,
            method="openpyxl"
        )
    except Exception as e:
        logger.error(f"XLSX extraction failed for {filename}: {e}")
        return MediaExtractionResult(
            source_type="file",
            filename=filename,
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            extracted_text="",
            success=False,
            error=f"XLSX extraction failed: {str(e)}",
            method="openpyxl"
        )


def extract_text_from_plaintext(file_bytes: bytes, filename: str = "file.txt", mime_type: str = "text/plain") -> MediaExtractionResult:
    """Extract text from plain text files (CSV, TXT, JSON, YAML, XML, source code, etc.)."""
    try:
        text = file_bytes.decode("utf-8", errors="replace").strip()
        return MediaExtractionResult(
            source_type="file",
            filename=filename,
            mime_type=mime_type,
            extracted_text=text,
            char_count=len(text),
            success=True,
            method="direct_read"
        )
    except Exception as e:
        logger.error(f"Text extraction failed for {filename}: {e}")
        return MediaExtractionResult(
            source_type="file",
            filename=filename,
            mime_type=mime_type,
            extracted_text="",
            success=False,
            error=f"Text extraction failed: {str(e)}",
            method="direct_read"
        )


# ── Main Dispatcher ──────────────────────────────────────────────────────────

def extract_from_attachment(
    content_base64: str,
    filename: str,
    mime_type: Optional[str] = None,
    max_size_bytes: int = 10 * 1024 * 1024
) -> MediaExtractionResult:
    """
    Extract text from a base64-encoded attachment.
    Routes to the appropriate extractor based on MIME type.
    """
    resolved_mime = _resolve_mime_type(filename, mime_type)

    try:
        file_bytes = _decode_base64_content(content_base64)
    except Exception as e:
        return MediaExtractionResult(
            source_type="file",
            filename=filename,
            mime_type=resolved_mime,
            extracted_text="",
            success=False,
            error=f"Base64 decode failed: {str(e)}",
            method="none"
        )

    # Check size limit
    if len(file_bytes) > max_size_bytes:
        return MediaExtractionResult(
            source_type="file",
            filename=filename,
            mime_type=resolved_mime,
            extracted_text="",
            success=False,
            error=f"File exceeds size limit ({len(file_bytes)} > {max_size_bytes} bytes)",
            method="none"
        )

    # Route to the appropriate extractor
    if resolved_mime in IMAGE_MIME_TYPES:
        result = extract_text_from_image(file_bytes, filename)
        result.mime_type = resolved_mime
        return result

    if resolved_mime == "application/pdf":
        return extract_text_from_pdf(file_bytes, filename)

    if resolved_mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return extract_text_from_docx(file_bytes, filename)

    if resolved_mime == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        return extract_text_from_xlsx(file_bytes, filename)

    # Fallback: try as plain text for text/* and JSON/XML/YAML
    if resolved_mime.startswith("text/") or resolved_mime in ("application/json", "application/xml", "application/x-yaml"):
        return extract_text_from_plaintext(file_bytes, filename, resolved_mime)

    # Unknown type — attempt plain text as last resort
    return extract_text_from_plaintext(file_bytes, filename, resolved_mime)


def extract_from_base64_image(data_uri_or_base64: str, filename: str = "inline_image") -> MediaExtractionResult:
    """
    Extract text from a base64-encoded image (e.g., from clipboard paste or inline embedding).
    Accepts both raw base64 and data URI format.
    """
    try:
        file_bytes = _decode_base64_content(data_uri_or_base64)
    except Exception as e:
        return MediaExtractionResult(
            source_type="image",
            filename=filename,
            mime_type="image/*",
            extracted_text="",
            success=False,
            error=f"Base64 decode failed: {str(e)}",
            method="none"
        )

    return extract_text_from_image(file_bytes, filename)


def process_all_media(
    attachments: Optional[List[dict]] = None,
    max_size_bytes: int = 10 * 1024 * 1024
) -> MediaScanSummary:
    """
    Process all media attachments and return a summary with combined extracted text.

    Each attachment dict should have:
      - filename: str
      - content_base64: str
      - mime_type: Optional[str]
    """
    summary = MediaScanSummary()

    if not attachments:
        return summary

    extracted_texts = []

    for attachment in attachments:
        filename = attachment.get("filename", "unknown")
        content_b64 = attachment.get("content_base64", "")
        mime = attachment.get("mime_type")

        if not content_b64:
            continue

        result = extract_from_attachment(content_b64, filename, mime, max_size_bytes)
        summary.extractions.append(result)
        summary.total_media_items += 1

        if result.source_type == "image":
            summary.images_scanned += 1
        else:
            summary.files_scanned += 1

        if result.success and result.extracted_text:
            extracted_texts.append(f"[Content from {filename}]:\n{result.extracted_text}")
            summary.total_text_extracted += result.char_count

    summary.combined_text = "\n\n".join(extracted_texts)
    return summary
